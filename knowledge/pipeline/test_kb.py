"""Self-check for kb.py — the gate's safety logic. Run: python test_kb.py

Plain asserts, no framework. Builds a throwaway wiki + staging tree in a temp
dir, points kb's module globals at it, and checks the two things that must hold:
validate catches broken proposals, and accept promotes only valid ones.
"""
import os
import tempfile

import kb


def _write(path, text):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)


def _stage(cycle_dir, relpath, text):
    _write(os.path.join(cycle_dir, "proposed", relpath), text)


GOOD_PROGRAM = """---
type: program
name: "Test Program"
status: Active
---
# Test Program
"""

# project links to a person staged in the SAME cycle -> link must still resolve
GOOD_PROJECT = """---
type: project
name: "Test Project"
status: Live
---
# Test Project
## People
- [[people/jane_doe/overview_jane_doe|Jane Doe]]
"""

GOOD_PERSON = """---
type: person
name: "Jane Doe"
status: Active
---
# Jane Doe
"""


def main():
    with tempfile.TemporaryDirectory() as tmp:
        # repoint kb at the sandbox
        kb.ROOT = tmp
        kb.STAGING = os.path.join(tmp, "_staging")
        kb.CHANGELOGS = os.path.join(tmp, "reporting", "changelogs")
        cycle = "2099-01-01"
        cdir = os.path.join(kb.STAGING, cycle)

        # 1. a clean, internally-consistent cycle validates
        _stage(cdir, "programs/test_program/overview_test_program.md", GOOD_PROGRAM)
        _stage(cdir, "projects/test_project/overview_test_project.md", GOOD_PROJECT)
        _stage(cdir, "people/jane_doe/overview_jane_doe.md", GOOD_PERSON)
        assert kb.validate(cycle) == [], "clean cycle should pass"

        # 2. missing required field is caught
        _stage(cdir, "people/bad/overview_bad.md",
               '---\ntype: person\nname: "Bad"\n---\n# Bad\n')  # no status
        errs = kb.validate(cycle)
        assert any("status" in m for _, m in errs), "missing status not caught"

        # 3. orphan wikilink is caught
        _stage(cdir, "projects/orphan/overview_orphan.md",
               '---\ntype: project\nname: "Orphan"\nstatus: Live\n---\n'
               '[[people/ghost/overview_ghost|Ghost]]\n')
        errs = kb.validate(cycle)
        assert any("orphan wikilink" in m for _, m in errs), "orphan link not caught"

        # 4. invalid YAML is caught
        _stage(cdir, "programs/broken/overview_broken.md",
               '---\ntype: program\nname: "x\nstatus: Active\n---\n# x\n')
        errs = kb.validate(cycle)
        assert any("YAML" in m or "frontmatter" in m for _, m in errs), "bad YAML not caught"

        # 5. accept refuses while the cycle is invalid (exit code 1)
        class A:
            pass
        a = A(); a.cycle = cycle; a.dry_run = False; a.only = None
        assert kb.cmd_accept(a) == 1, "accept must refuse an invalid cycle"

        # 6. a valid single-file cycle promotes into the wiki + writes changelog
        good_cycle = "2099-02-02"
        gdir = os.path.join(kb.STAGING, good_cycle)
        _stage(gdir, "people/jane_doe/overview_jane_doe.md", GOOD_PERSON)
        a2 = A(); a2.cycle = good_cycle; a2.dry_run = False; a2.only = None
        assert kb.cmd_accept(a2) == 0, "valid cycle should accept"
        promoted = os.path.join(tmp, "people", "jane_doe", "overview_jane_doe.md")
        assert os.path.isfile(promoted), "file not promoted into wiki"
        log = os.path.join(kb.CHANGELOGS, f"{good_cycle}.md")
        assert os.path.isfile(log), "changelog not written"

        # 7. seed: a structured workbook stages an internally-valid baseline
        import openpyxl
        xlsx = os.path.join(tmp, "roster.xlsx")
        wbk = openpyxl.Workbook()
        ws = wbk.active; ws.title = "people"
        ws.append(["name", "aliases", "role", "leader", "status", "programs", "projects", "tag"])
        ws.append(["Jane Doe", "Janie; JD", "DS", "Matt", "Active",
                   "rev_mgmt", "rev_mgmt/track_opt", "jane_doe"])
        ws = wbk.create_sheet("programs")
        ws.append(["program", "aliases", "status", "goal", "type", "tag"])
        ws.append(["Rev Mgmt", "RMA", "Active", "x", "performance", "rev_mgmt"])
        ws = wbk.create_sheet("projects")
        ws.append(["project", "aliases", "program", "status", "summary", "tag"])
        ws.append(["Track Opt", "TO", "Rev Mgmt", "In Development", "y", "rev_mgmt/track_opt"])
        wbk.save(xlsx)

        s = A(); s.file = xlsx; s.template = None; s.cycle = "seed-test"
        assert kb.cmd_seed(s) == 0, "seed should stage a valid sheet"
        assert kb.validate("seed-test") == [], "seeded baseline should validate"
        prop = os.path.join(kb.STAGING, "seed-test", "proposed")
        ptext = open(os.path.join(prop, "people", "jane_doe", "overview_jane_doe.md"), encoding="utf-8").read()
        assert 'aliases: ["Janie", "JD"]' in ptext and "tag: jane_doe" in ptext, "person tag/alias"
        gtext = open(os.path.join(prop, "programs", "rev_mgmt", "overview_rev_mgmt.md"), encoding="utf-8").read()
        assert "RMA" in gtext and "category: performance" in gtext, "program alias/category"
        # project nests under its program tag: projects/<prog>/<proj>/
        jtext = open(os.path.join(prop, "projects", "rev_mgmt", "track_opt",
                                  "overview_track_opt.md"), encoding="utf-8").read()
        assert "TO" in jtext and "tag: rev_mgmt/track_opt" in jtext, "project tag/alias"
        idx = open(os.path.join(prop, "indexes", "programs.md"), encoding="utf-8").read()
        assert "(aka: RMA)" in idx, "alias not surfaced in index lookup table"

        # 8. structural hard-fail: a project whose program column is unknown
        bad = os.path.join(tmp, "bad.xlsx")
        wb2 = openpyxl.Workbook()
        ws = wb2.active; ws.title = "people"
        ws.append(["name", "aliases", "role", "leader", "status", "programs", "projects", "tag"])
        ws = wb2.create_sheet("programs")
        ws.append(["program", "aliases", "status", "goal", "type", "tag"]); ws.append(["P", "", "Active", "g", "", "p"])
        ws = wb2.create_sheet("projects")
        ws.append(["project", "aliases", "program", "status", "summary", "tag"])
        ws.append(["Orphan Proj", "", "Ghost Program", "Planned", "s", ""])
        wb2.save(bad)
        s2 = A(); s2.file = bad; s2.template = None; s2.cycle = "seed-bad"
        assert kb.cmd_seed(s2) == 1, "seed must reject a project with an unknown program"

        # 9. a person referencing an unknown program is a warning, not a failure
        soft = os.path.join(tmp, "soft.xlsx")
        wb3 = openpyxl.Workbook()
        ws = wb3.active; ws.title = "people"
        ws.append(["name", "aliases", "role", "leader", "status", "programs", "projects", "tag"])
        ws.append(["Cy", "", "DS", "M", "Active", "nonexistent_program", "", "cy"])
        ws = wb3.create_sheet("programs")
        ws.append(["program", "aliases", "status", "goal", "type", "tag"]); ws.append(["P", "", "Active", "g", "", "p"])
        ws = wb3.create_sheet("projects")
        ws.append(["project", "aliases", "program", "status", "summary", "tag"])
        wb3.save(soft)
        s3 = A(); s3.file = soft; s3.template = None; s3.cycle = "seed-soft"
        assert kb.cmd_seed(s3) == 0, "unresolved person ref should warn, not block"
        assert kb.validate("seed-soft") == [], "soft-seed baseline should still validate"

    print("OK - all kb.py gate checks passed")


if __name__ == "__main__":
    main()
