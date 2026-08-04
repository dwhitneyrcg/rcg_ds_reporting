from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# ============================================================
# SECTION 1: WEEKLY HIGHLIGHTS (EVP / Jason Update)
# ============================================================
h = doc.add_heading('Weekly AI Highlights — April 17, 2026', level=1)
h.runs[0].font.size = Pt(16)

# --- Achievements ---
h2 = doc.add_heading('This Week — 3 Key AI Achievements', level=2)
h2.runs[0].font.size = Pt(13)

achievements = [
    (
        "RCI Pricing Automation — Leadership Alignment on Recommendation Accuracy",
        "Analysis presented to senior leadership demonstrated that 30% of seemingly counterintuitive PRE recommendations are actually correct, driven by demand forecast–track misalignment rather than model error. This pivotal change-management moment is expected to significantly increase PRE adoption by product teams, with track smoothing delivery targeted for early May."
    ),
    (
        "Supply Chain AI Pilot — Business Endorsement for Fleetwide Expansion",
        "The Beyond pilot convinced business stakeholders that AI demand forecasting outperforms rolling-average baselines, with conversations now shifting from pilot validation to fleetwide adoption—a major milestone for the $14M supply chain cost-reduction initiative."
    ),
    (
        "MyCruise Recommendation Engine — 96% Coverage with 4x Faster Training",
        "Delivered an FP Growth model that raised recommendation coverage from ~60% to ~96%, while Bayesian optimization cut ALS training time from 20 minutes to under 5 minutes per cycle. Apriori fallback covers 80–83% of empty collaborative-filtering results with only ~5% overlap, surfacing genuinely new products."
    ),
]

for headline, sentence in achievements:
    p = doc.add_paragraph()
    run_hl = p.add_run(headline)
    run_hl.bold = True
    run_hl.font.size = Pt(11)
    p.add_run('\n' + sentence).font.size = Pt(11)

# --- Focus Areas ---
h2 = doc.add_heading('Near Term — 4 Focus Areas', level=2)
h2.runs[0].font.size = Pt(13)

focus_areas = [
    (
        "CEL Track Optimization — Hierarchical Smoothing for Optimal Pricing Targets",
        "A new two-stage approach combining dynamic programming with MILP smoothing produces near-optimal weekly track asks in 27 seconds, eliminating the noisy week-over-week fluctuations that misalign with business operations. Results for both RCI and CEL will be presented to stakeholders next week."
    ),
    (
        "International Call Volume Forecasting — 16-Market Scalable Pipeline",
        "A 130-feature forecasting pipeline with automated market-stability tiering is ready for its first end-to-end production run across all International and Casino markets, directly enabling workforce planning optimization."
    ),
    (
        "PCP Mass Promo and 1:1 Targeted Offer Interfaces",
        "Prototype UIs for both mass promotion management and individualized targeted offers are in active development, positioning the team to deliver end-to-end automated pricing and promotion workflows."
    ),
    (
        "Co-Brand Loyalty Profile Merge — 25,300 Records Staged for Launch",
        "LLM-validated profile merges with cosine similarity >0.8 and same-program deduplication rules are staged for execution, supporting the Co-Brand launch with cleaner, consolidated guest profiles."
    ),
]

for headline, sentence in focus_areas:
    p = doc.add_paragraph()
    run_hl = p.add_run(headline)
    run_hl.bold = True
    run_hl.font.size = Pt(11)
    p.add_run('\n' + sentence).font.size = Pt(11)

# Page break before detailed report
doc.add_page_break()

# ============================================================
# SECTION 2: DETAILED BUSINESS AREA UPDATES
# ============================================================
h = doc.add_heading('Business Area Updates — April 17, 2026', level=1)
h.runs[0].font.size = Pt(16)

updates = [
    (
        "AXIOM: Generative-AI SEATs (Hotel Ops, Consumer Insights, Legal)",
        "Deployed a port report into a Port Operations-facing dashboard pipeline and completed a Gratuity & Tips email report now pending business review. Prepared a presentation for Hotel Operations on critical data cleaning requirements for guest log classifications—work that will unlock trend tracking, summary generation, and predictive modeling across Voice 360.",
        "Teams are building a validation dashboard comparing AI, naive, and business targets for customer satisfaction forecasting, migrating the Axiom web app to Azure Container Apps, and awaiting stakeholder input on guest logs gold table discrepancies."
    ),
    (
        "Contact Center Optimization & Automation (RCI/CEL)",
        "Delivered a scalable call-volume forecasting pipeline for International and Casino markets spanning 16+ geographies, featuring ~130 engineered features adapted from a validated supply-chain forecasting framework, synthetic forecast baselines for future periods, and adaptive configuration that automatically tunes models by market stability tier.",
        "Teams are validating forecasting functions and preparing the first end-to-end forecast run across all markets."
    ),
    (
        "Customer Lifetime Value (Corporate Planning)",
        "Delivered hurricane risk analysis findings to the AVP of Corporate Strategy, providing quantitative inputs to inform capital allocation decisions for private destination investments. Resolved passenger duplication defects and row count inflation in the CLV table pipeline caused by faulty join logic against outdated source tables, and communicated reproducible defect examples to upstream data owners for remediation.",
        "Teams are conducting strategic analysis of MSC's private destination expansion across the Caribbean and validating metrics within executive-level deliverables using the competitive cruise fleet model."
    ),
    (
        "Customer Targeting (E-Commerce)",
        "Completed backtesting for consumer segmentation and targeted offer strategy, revealing that booking propensity shows a clear linear relationship with actual booking rates while web propensity does not meaningfully differentiate booking behavior—findings that may lead to refinement of the segmentation logic for more effective targeting. Implemented new app-usage-driven features into the booking propensity model (post-voyage usage rate, onboard usage count, engagement recency signals) and expanded over 20 web-related features for the journey scoring model.",
        "Teams are aligning with business stakeholders to reassess segmentation strategy and potentially adjust the role of web propensity in targeting criteria."
    ),
    (
        "Hybris Product Recommendations (Digital)",
        "Shipped the guest segmentation ETL migrated to PySpark for scale with a data-leakage fix, and achieved a 4x ALS collaborative filtering training speedup (20 min to 5 min) through Bayesian hyperparameter optimization. Delivered an FP Growth model achieving ~96% recommendation coverage—up from ~60%—with Apriori fallback covering 80–83% when collaborative filtering returns empty and only ~5% overlap, surfacing genuinely new products. Dev Postgres load time improved from 10+ minutes to under 4 minutes after infrastructure upgrade.",
        "Teams are validating product clustering embeddings against a 35K-product dataset, implementing hard-coded business override slots in the recommendations API, and awaiting calendar schema deployment to unblock the ForYou model coverage fix."
    ),
    (
        "Loyalty Program Redesign",
        "Recommended 25,300 LLM-validated profile merges for the Co-Brand launch, applying cosine similarity thresholds (>0.8), same-program deduplication rules, and single Consumer ID constraints to ensure merge quality and prevent erroneous profile consolidation.",
        "Teams are executing the profile merge deployment and monitoring results."
    ),
    (
        "Marine Insights Analytics Platform (Marine Operations)",
        "Met with Valmet at SeaTrade to resolve ongoing OPC UA server issues and reviewed Lufthansa's Digital Twin solution for the cruise industry, scheduling a detailed follow-up with Data Science teams. Created fleet-wide Open-Meteo weather tables enabling digital twin predictions for ships without onboard weather data, productionized the Residual Boosting class into the digital twin package, and presented 2025 fuel forecast results versus actuals and finance predictions to Decarb Digital Solutions stakeholders.",
        "Teams are completing the MIAP app refactoring to Next.js/FastAPI for improved performance and agent integration, improving SFOC curves, developing individual chiller models, and investigating a 200kW AHU deviation and ERW usage drop with Marine Engineering."
    ),
    (
        "PCP Pricing Automation (RCI/CEL)",
        "Delivered Flash Sale casino member impact analysis showing positive revenue lift for alcoholic passes (PRIME loyalty tier highest increase) across most meta products except Short Caribbean, and advanced CEL Beverage EDA Phase 2 with clear elasticity relationships identified at the meta product, ship class, and weeks-to-sail levels. Developed an initial pricing recommendation framework for RBC Passes and Waterpark products using four booking phase segments and a composite sailing performance score, and refined waterpark A/B test pairing with tighter constraints (7-day window, same day-of-week, same ship).",
        "Teams are building prototype UIs for mass promotions and 1:1 targeted offers, testing enhanced promo automation features in lower environments, and establishing an automated product classification pipeline using multi-pass LLM consensus methodology."
    ),
    (
        "Revenue Management Automation (CEL)",
        "Reviewed Category-Gapping 3.0 optimization results with business stakeholders and built eight availability-dependent pricing formulas with a unified process that evaluates tier eligibility, expected value, booking share, and inventory limits to select revenue-maximizing options across all future sailings. Completed promotions analysis revealing that replacement-value promos are driven primarily by remaining inventory rather than track performance, with the strongest effect in the 0–12 weeks-to-sail window—Short Caribbean shows high replacement reliance while Europe/Alaska shows low dependence, informing more targeted promo deployment.",
        "Teams are developing a hierarchical track optimization combining dynamic programming with MILP smoothing (27-second runtime on CEL Europe), analyzing SPI-based top performer tracks, and presenting results to stakeholders next week."
    ),
    (
        "Revenue Management Automation (RCI)",
        "Presented PRE analysis to senior leadership demonstrating that 30% of counterintuitive pricing recommendations stem from demand forecast–future track disconnects rather than model error—a pivotal change-management moment expected to drive significantly higher PRE adoption by product teams, with action items including track smoothing (delivery by early May) and enhanced reporting for paradoxical pricing situations. Completed the metaproduct basket framework for Short Caribbean, Alaska, Europe, Galapagos, Mexico, and West Coast Short, standardizing booking pace feature engineering with statistical peer group assignment and pricing signal classification across five action categories (~90% complete).",
        "Teams are extending basket development to Long Caribbean, training Category-Gapping 3.0 EBM models at multiple granularity levels, finalizing the GTY booking volume KPI quantification model, building demand model ensembles with external data integration (Google Search Trends, FRED macro indicators), and transitioning SPI Portal features to the dev environment."
    ),
    (
        "Supply Chain Optimization",
        "Delivered improved voyage forecast prorations to production and provided Beyond shipboard inventory managers with AI demand model visualizations demonstrating superiority over simple rolling averages for products with downward consumption trends—business stakeholders are now shifting conversations from the BY pilot toward fleetwide adoption, a major milestone. Fixed Predicted Voyage Demand drift, eliminated daily requisition ghost rows doubling per-guest variance, replaced the legacy consumption ratio approach with a hybrid voyage-lagged method, and resolved SSC Finance Tool issues including SHIP_CODE derivation, 2x cost doubling, and historical drift.",
        "Teams are refactoring the finance tool into 14 smaller notebooks, testing SSC crew count model impact on downstream predictions, hardening ETL bid data pipelines, and preparing for PRD GOLD decimal data-type regression testing."
    ),
    (
        "Win-on-Waste (Hotel Operations)",
        None,
        "Teams are focused on resolving bugs in the inter-port and specialty dining V2 forecasting modules."
    ),
]

for biz_area, accomplishments, ongoing in updates:
    p = doc.add_paragraph(style='List Bullet')
    run_area = p.add_run(biz_area + ': ')
    run_area.bold = True
    run_area.font.size = Pt(11)
    if accomplishments:
        run_acc = p.add_run(accomplishments + ' ')
        run_acc.font.size = Pt(11)
    if ongoing:
        run_ong = p.add_run(ongoing)
        run_ong.italic = True
        run_ong.font.size = Pt(11)

# Save
output_path = r'C:\Users\133486\OneDrive - Royal Caribbean Group\20260417 - Weekly Matt & Rafeh Update.docx'
doc.save(output_path)
print(f'Report saved to: {output_path}')
