"""Convert .docx files to markdown, preserving text formatting."""
import sys
import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE


def run_to_md(run):
    """Convert a single run to markdown with inline formatting."""
    text = run.text
    if not text:
        return ""
    if run.bold and run.italic:
        text = f"***{text}***"
    elif run.bold:
        text = f"**{text}**"
    elif run.italic:
        text = f"*{text}*"
    if run.underline:
        text = f"<u>{text}</u>"
    if run.font.strike:
        text = f"~~{text}~~"
    if run.font.superscript:
        text = f"<sup>{text}</sup>"
    if run.font.subscript:
        text = f"<sub>{text}</sub>"
    return text


def get_heading_level(para):
    """Return heading level (1-9) or 0 if not a heading."""
    style_name = para.style.name if para.style else ""
    if style_name.startswith("Heading"):
        try:
            return int(style_name.split()[-1])
        except ValueError:
            pass
    if style_name == "Title":
        return 1
    if style_name == "Subtitle":
        return 2
    return 0


def get_list_info(para):
    """Return (list_level, is_numbered) or None if not a list."""
    style_name = para.style.name if para.style else ""
    # Check numbering via XML
    pPr = para._element.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr')
    if pPr is not None:
        ilvl = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ilvl')
        numId = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId')
        level = int(ilvl.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '0')) if ilvl is not None else 0
        # Heuristic: check style name for numbered vs bullet
        is_numbered = "list number" in style_name.lower() or "numbered" in style_name.lower()
        return (level, is_numbered)
    if "list bullet" in style_name.lower():
        return (0, False)
    if "list number" in style_name.lower():
        return (0, True)
    return None


def para_to_md(para):
    """Convert a paragraph to markdown."""
    # Assemble inline text
    text = "".join(run_to_md(r) for r in para.runs)
    if not text.strip():
        return ""

    # Headings
    level = get_heading_level(para)
    if level:
        return f"{'#' * level} {text.strip()}"

    # Lists
    list_info = get_list_info(para)
    if list_info is not None:
        indent = "  " * list_info[0]  # nested indent
        if list_info[1]:
            return f"{indent}1. {text.strip()}"
        else:
            return f"{indent}- {text.strip()}"

    return text.strip()


def table_to_md(table):
    """Convert a table to markdown."""
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            cell_text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            # Escape pipes
            cell_text = cell_text.replace("|", "\\|")
            cells.append(cell_text)
        rows.append(cells)

    if not rows:
        return ""

    lines = []
    # Header row
    lines.append("| " + " | ".join(rows[0]) + " |")
    lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
    for row in rows[1:]:
        # Pad if needed
        while len(row) < len(rows[0]):
            row.append("")
        lines.append("| " + " | ".join(row[:len(rows[0])]) + " |")
    return "\n".join(lines)


def convert_docx_to_md(docx_path):
    """Convert a .docx file to markdown string."""
    doc = Document(docx_path)
    md_parts = []

    # Iterate over body elements in order (paragraphs and tables)
    from docx.oxml.ns import qn
    body = doc.element.body
    for child in body:
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag == 'p':
            # Find the matching paragraph object
            for para in doc.paragraphs:
                if para._element is child:
                    line = para_to_md(para)
                    md_parts.append(line)
                    break
        elif tag == 'tbl':
            for table in doc.tables:
                if table._element is child:
                    md_parts.append("")
                    md_parts.append(table_to_md(table))
                    md_parts.append("")
                    break

    # Clean up multiple blank lines
    result = "\n\n".join(part for part in md_parts if part is not None)
    # Collapse 3+ newlines to 2
    import re
    result = re.sub(r'\n{3,}', '\n\n', result)
    return result.strip() + "\n"


def main():
    import argparse

    parser = argparse.ArgumentParser(description="Convert .docx files to Markdown.")
    parser.add_argument(
        "input_dir",
        nargs="?",
        default=r"C:\Users\133486\OneDrive - Royal Caribbean Group\Co-Pilot Demo\Strat Plan Demo\raw_inputs",
        help="Directory containing .docx files (default: Strat Plan Demo/raw_inputs)",
    )
    parser.add_argument(
        "-o", "--output-dir",
        default=None,
        help="Output directory for .md files (default: sibling 'markdown_inputs' folder)",
    )
    args = parser.parse_args()

    src_dir = args.input_dir
    if args.output_dir:
        out_dir = args.output_dir
    else:
        out_dir = os.path.join(os.path.dirname(src_dir), "markdown_inputs")

    os.makedirs(out_dir, exist_ok=True)

    # Find all .docx files in src_dir
    files = [f for f in os.listdir(src_dir) if f.lower().endswith(".docx") and not f.startswith("~$")]
    if not files:
        print(f"No .docx files found in {src_dir}")
        return

    files.sort()
    print(f"Source: {src_dir}")
    print(f"Output: {out_dir}")
    print(f"Found {len(files)} .docx file(s)\n")

    for f in files:
        docx_path = os.path.join(src_dir, f)
        md_name = os.path.splitext(f)[0] + ".md"
        md_path = os.path.join(out_dir, md_name)

        print(f"Converting: {f}")
        md_content = convert_docx_to_md(docx_path)
        with open(md_path, "w", encoding="utf-8") as fp:
            fp.write(md_content)
        print(f"  -> {md_name} ({len(md_content):,} chars)")

    print(f"\nDone. {len(files)} file(s) converted.")


if __name__ == "__main__":
    main()
